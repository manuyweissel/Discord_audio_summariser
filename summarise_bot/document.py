from __future__ import annotations

import re
from datetime import datetime
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from .logger import logger


def extract_meeting_title(summary: str, default: str = "Meeting") -> str:
    """First non-empty line under the '## Meeting Topic' heading — used as the doc subtitle."""
    lines = summary.splitlines()
    for index, line in enumerate(lines):
        if line.strip().lower().startswith("## meeting topic"):
            for following in lines[index + 1:]:
                text = following.strip()
                if text:
                    return text
            break
    return default


_BOLD_RE = re.compile(r"\*\*(.+?)\*\*", re.S)
_RULE_RE = re.compile(r"^\s*([-*_])\1{2,}\s*$")


def extract_meeting_date(summary: str) -> str | None:
    """The meeting date from the '- **Date:** dd.mm.yyyy' bullet the summarizer injects.

    Validated against a date shape so an unrelated bullet can never end up in the subtitle.
    """
    for line in summary.splitlines():
        match = re.match(r"^\s*[-*+]\s+\*{0,2}\s*Date:\*{0,2}\s*(.+?)\s*$", line, re.I)
        if match:
            value = match.group(1).strip()
            if re.fullmatch(r"\d{1,2}[./-]\d{1,2}[./-]\d{2,4}", value):
                return value
    return None


def _add_rich_text(paragraph, text: str) -> None:
    """Render **bold** spans as real runs instead of literal asterisks."""
    cursor = 0
    for match in _BOLD_RE.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        paragraph.add_run(match.group(1)).bold = True
        cursor = match.end()
    remainder = text[cursor:]
    if remainder or cursor == 0:
        paragraph.add_run(remainder)


def convert_to_word_doc(
    content: str, meeting_title: str = "Meeting", meeting_date: str | None = None
) -> bytes | None:
    try:
        doc = Document()
        title = doc.add_paragraph()
        title.style = doc.styles["Title"]
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        title_run = title.add_run("Meeting Minutes")
        title_run.font.size = Pt(22)
        # Fall back to today only when the summary carries no meeting date; the docx used to
        # always show its own generation date rather than when the meeting happened.
        subtitle_date = meeting_date or extract_meeting_date(content) or datetime.now().strftime("%d.%m.%Y")
        subtitle = doc.add_paragraph(f"{meeting_title} • {subtitle_date}")
        subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT

        table_rows: list[list[str]] = []
        in_table = False
        for raw_line in content.splitlines():
            line = raw_line.strip()
            if not line:
                if in_table and table_rows:
                    _append_table(doc, table_rows)
                    table_rows = []
                    in_table = False
                continue

            if _RULE_RE.match(line):
                # Markdown horizontal rule: structural, not content. Previously printed as "---".
                continue

            if line.startswith("# ") and not line.startswith("## "):
                if in_table and table_rows:
                    _append_table(doc, table_rows)
                    table_rows = []
                    in_table = False
                # The document already carries its own "Meeting Minutes" title block, so the
                # markdown H1 would duplicate it. Previously it leaked through as body text
                # complete with the hash character.
                continue

            if line.startswith("## "):
                if in_table and table_rows:
                    _append_table(doc, table_rows)
                    table_rows = []
                    in_table = False
                doc.add_heading(line.replace("## ", "").strip(), level=1)
                continue

            if line.startswith("### "):
                if in_table and table_rows:
                    _append_table(doc, table_rows)
                    table_rows = []
                    in_table = False
                doc.add_heading(line.replace("### ", "").strip(), level=2)
                continue

            if "|" in line:
                parts = line.split("|")
                # Drop only the empty artifacts from the leading/trailing pipes; keep interior
                # cells even when blank so columns (e.g. an unstated Owner/Due) stay aligned.
                if parts and parts[0].strip() == "":
                    parts = parts[1:]
                if parts and parts[-1].strip() == "":
                    parts = parts[:-1]
                cells = [cell.strip() for cell in parts]
                # Skip the markdown separator/alignment row (cells made only of '-' and ':',
                # e.g. '---', ':--:', '---:').
                is_separator = bool(cells) and all(cell and set(cell) <= {"-", ":"} for cell in cells)
                if cells and not is_separator:
                    table_rows.append(cells)
                    in_table = True
                continue

            paragraph = doc.add_paragraph()
            if line.startswith(("- ", "* ")):
                paragraph.style = "List Bullet"
                _add_rich_text(paragraph, line[2:].strip())
            else:
                _add_rich_text(paragraph, line)

        if in_table and table_rows:
            _append_table(doc, table_rows)

        footer = doc.add_paragraph()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run(f"Generated automatically • {datetime.now().strftime('%d.%m.%Y %H:%M')}")

        output = BytesIO()
        doc.save(output)
        return output.getvalue()
    except Exception as error:
        logger.error(
            "Failed to render meeting minutes document",
            extra={"action": "document_render", "event": "error", "error_message": str(error)},
        )
        return None


def _append_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=1, cols=width)
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for idx, value in enumerate(rows[0]):
        header_cells[idx].text = value
    for row_values in rows[1:]:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cells[idx].text = value
