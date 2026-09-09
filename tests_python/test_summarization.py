from __future__ import annotations

from datetime import datetime
from io import BytesIO
import unittest

from docx import Document as DocxDocument

from summarise_bot.document import convert_to_word_doc, extract_meeting_title
from summarise_bot.summarization import _apply_known_facts, _extract_meeting_facts


SAMPLE_TRANSCRIPT = (
    "[2026-06-11T13:05:36.500000+00:00] Manú: Los geht's.\n"
    "[2026-06-11T13:06:10.000000+00:00] Mark: Yep.\n"
    "[2026-06-11T13:50:00.000000+00:00] Minh Tú: Bis dann.\n"
)


class ExtractMeetingFactsTests(unittest.TestCase):
    def test_participants_in_first_seen_order(self) -> None:
        facts = _extract_meeting_facts(SAMPLE_TRANSCRIPT)
        self.assertEqual(facts["participants"], "Manú, Mark, Minh Tú")

    def test_date_and_time_local_and_formatted(self) -> None:
        facts = _extract_meeting_facts(SAMPLE_TRANSCRIPT)
        expected_date = (
            datetime.fromisoformat("2026-06-11T13:05:36.500000+00:00").astimezone().strftime("%d.%m.%Y")
        )
        self.assertEqual(facts["date"], expected_date)
        self.assertRegex(facts["time"], r"^\d{2}:\d{2}–\d{2}:\d{2}$")

    def test_no_facts_when_lines_unparseable(self) -> None:
        self.assertEqual(_extract_meeting_facts("just some text without prefixes"), {})


class ApplyKnownFactsTests(unittest.TestCase):
    BLUEPRINT = "# Meeting Minutes — {{DATE}}\n\n## Details\n{{DETAILS}}\n"

    def test_fills_header_with_all_facts(self) -> None:
        out = _apply_known_facts(
            self.BLUEPRINT,
            {"date": "11.06.2026", "time": "15:05–15:50", "participants": "A, B"},
        )
        self.assertIn("# Meeting Minutes — 11.06.2026", out)
        self.assertIn("- **Date:** 11.06.2026", out)
        self.assertIn("- **Time:** 15:05–15:50", out)
        self.assertIn("- **Participants:** A, B", out)
        self.assertNotIn("{{", out)

    def test_omits_missing_lines_and_defaults_date(self) -> None:
        out = _apply_known_facts(self.BLUEPRINT, {})
        self.assertNotIn("Time:", out)
        self.assertNotIn("Participants:", out)
        self.assertNotIn("n/a", out)
        self.assertNotIn("{{", out)


class ExtractMeetingTitleTests(unittest.TestCase):
    def test_returns_topic_not_date(self) -> None:
        summary = (
            "# Meeting Minutes — 11.06.2026\n\n"
            "## Meeting Topic\nProduct Launch Strategy\n\n"
            "## Details\n- **Date:** 11.06.2026\n"
        )
        self.assertEqual(extract_meeting_title(summary), "Product Launch Strategy")

    def test_default_when_no_topic(self) -> None:
        self.assertEqual(extract_meeting_title("# Meeting Minutes — 11.06.2026\n"), "Meeting")


class DocxTableTests(unittest.TestCase):
    def test_table_skips_separator_and_preserves_blank_cells(self) -> None:
        markdown = (
            "## Action Items\n"
            "| # | Task | Owner | Due |\n"
            "|---:|---|---|---|\n"
            "| 1 | Do the thing | Alice |  |\n"
            "| 2 | Other thing |  |  |\n"
        )
        data = convert_to_word_doc(markdown, "Test")
        self.assertIsNotNone(data)
        doc = DocxDocument(BytesIO(data))
        self.assertEqual(len(doc.tables), 1)
        table = doc.tables[0]
        self.assertEqual(len(table.columns), 4)
        self.assertEqual(len(table.rows), 3)  # header + 2 data rows, separator skipped
        self.assertEqual([c.text for c in table.rows[0].cells], ["#", "Task", "Owner", "Due"])
        self.assertEqual([c.text for c in table.rows[2].cells], ["2", "Other thing", "", ""])


if __name__ == "__main__":
    unittest.main()
